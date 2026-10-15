"""Tests for the internship calendar: day-by-day calculation, overrides and XLSX export."""

import json
from datetime import date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models import AcademicYear, ExcludedPeriod
from app.routes.internship import InternshipRequest, _export_filename
from app.services.internship import (
    DAY_KIND_HOLIDAY,
    DAY_KIND_OVERRIDE,
    DAY_KIND_WEEKEND,
    DAY_KIND_WORKING,
    InternshipSchedule,
    build_internship_calendar,
    internship_deadline,
)
from app.services.internship_export import METADATA_FIELDS, build_internship_workbook
from app.services.internship_html import CALENDAR_KIND_LABELS, build_month_views, inline_stylesheet

WEEKDAYS_7H = {0: 420, 1: 420, 2: 420, 3: 420, 4: 420}


def _calendar(target_minutes, weekday_minutes=None, non_working=None, overrides=None, start=date(2026, 10, 12)):
    # 12/10/2026 is a Monday.
    return build_internship_calendar(start, target_minutes, weekday_minutes or WEEKDAYS_7H, non_working or set(), overrides)


def test_end_date_counts_from_start_day_and_skips_weekends():
    """35 h at 7 h/day fills Monday to Friday of the first week."""

    result = _calendar(35 * 60)
    assert result.end_date == date(2026, 10, 16)
    assert result.working_day_count == 5

    result = _calendar(42 * 60)
    # Saturday and Sunday contribute nothing, so the sixth day is the next Monday.
    assert result.end_date == date(2026, 10, 19)
    assert [day.kind for day in result.days[5:7]] == [DAY_KIND_WEEKEND, DAY_KIND_WEEKEND]


def test_holidays_are_skipped():
    holiday = date(2026, 10, 14)
    result = _calendar(35 * 60, non_working={holiday})
    assert result.end_date == date(2026, 10, 19)
    holiday_day = next(day for day in result.days if day.date == holiday)
    assert holiday_day.minutes == 0
    assert holiday_day.kind == DAY_KIND_HOLIDAY


def test_minutes_in_daily_schedule_are_exact():
    """4 h 30 min per day: 9 h ends on the second day without rounding."""

    result = _calendar(9 * 60, weekday_minutes={index: 270 for index in range(5)})
    assert result.end_date == date(2026, 10, 13)
    assert result.scheduled_minutes == 540


def test_last_day_only_gets_remaining_minutes():
    result = _calendar(10 * 60)
    last = result.days[-1]
    assert last.date == date(2026, 10, 13)
    assert last.minutes == 180
    assert last.is_partial
    assert result.scheduled_minutes == 600


def test_override_reducing_hours_moves_end_date_later():
    result = _calendar(35 * 60, overrides={date(2026, 10, 13): 0})
    assert result.end_date == date(2026, 10, 19)
    assert result.days[1].kind == DAY_KIND_OVERRIDE
    assert result.days[1].minutes == 0


def test_override_adding_hours_on_weekend_moves_end_date_earlier():
    base = _calendar(42 * 60)
    assert base.end_date == date(2026, 10, 19)
    result = _calendar(42 * 60, overrides={date(2026, 10, 17): 420})
    assert result.end_date == date(2026, 10, 17)
    saturday = result.days[5]
    assert saturday.kind == DAY_KIND_OVERRIDE and saturday.minutes == 420


def test_override_on_holiday_counts_as_working_time():
    holiday = date(2026, 10, 14)
    result = _calendar(35 * 60, non_working={holiday}, overrides={holiday: 420})
    assert result.end_date == date(2026, 10, 16)


def test_override_matching_default_is_not_flagged():
    result = _calendar(35 * 60, overrides={date(2026, 10, 13): 420})
    assert result.days[1].kind == DAY_KIND_WORKING


def test_zero_schedule_is_rejected():
    with pytest.raises(ValueError):
        _calendar(10 * 60, weekday_minutes={index: 0 for index in range(5)})


def test_deadline_is_31_july_of_the_school_year():
    assert internship_deadline(date(2026, 9, 1)) == date(2027, 7, 31)
    assert internship_deadline(date(2026, 12, 31)) == date(2027, 7, 31)
    assert internship_deadline(date(2027, 3, 1)) == date(2027, 7, 31)
    assert internship_deadline(date(2027, 7, 31)) == date(2027, 7, 31)
    with pytest.raises(ValueError, match="agost"):
        internship_deadline(date(2027, 8, 2))


def test_hours_that_do_not_fit_before_august_stop_on_31_july():
    """10 h/week from June cannot reach 200 h: the calendar stops on 31 July with hours pending."""

    result = _calendar(200 * 60, weekday_minutes={index: 120 for index in range(5)}, start=date(2027, 6, 1))
    assert result.end_date == date(2027, 7, 31)
    assert all(day.date.month != 8 for day in result.days)
    assert not result.is_complete
    assert result.last_working_date == date(2027, 7, 30)  # Friday; 31 July is a Saturday.
    # June 2027 has 22 weekdays and July 22: 44 days of 2 h.
    assert result.scheduled_minutes == 44 * 120
    assert result.pending_minutes == 200 * 60 - 44 * 120


def test_override_minutes_alone_stop_at_the_deadline():
    result = _calendar(100 * 60, weekday_minutes={index: 0 for index in range(5)}, overrides={date(2026, 10, 13): 60})
    assert result.end_date == date(2027, 7, 31)
    assert result.scheduled_minutes == 60
    assert result.pending_minutes == 99 * 60


def test_completed_internship_has_no_pending_minutes():
    result = _calendar(35 * 60)
    assert result.is_complete and result.pending_minutes == 0
    assert result.last_working_date == result.end_date


def _period(period_id, start, end, label):
    return ExcludedPeriod(id=period_id, start_date=start, end_date=end, label=label)


def _payload(**overrides):
    payload = {
        "start_date": "12/10/2026",
        "target_hours": "35",
        "target_minutes": 0,
        "weekday_minutes": {str(index): 420 for index in range(5)},
        "excluded_period_ids": [],
        "extra_periods": [],
        "overrides": {},
        "metadata": {"student_name": "Núria Pérez"},
    }
    payload.update(overrides)
    return payload


def test_request_applies_only_selected_school_periods_and_extra_dates():
    periods = [
        _period(1, date(2026, 10, 13), date(2026, 10, 13), "Festa A"),
        _period(2, date(2026, 10, 14), date(2026, 10, 14), "Festa B"),
    ]
    request = InternshipRequest(
        _payload(excluded_period_ids=[1], extra_periods=[{"start": "15/10/2026", "end": "", "label": "Empresa tancada"}]),
        periods,
    )
    result = request.calculate()
    kinds = {day.date: (day.kind, day.label) for day in result.days}
    assert kinds[date(2026, 10, 13)] == (DAY_KIND_HOLIDAY, "Festa A")
    assert kinds[date(2026, 10, 14)] == (DAY_KIND_WORKING, "")
    assert kinds[date(2026, 10, 15)] == (DAY_KIND_HOLIDAY, "Empresa tancada")
    assert result.end_date == date(2026, 10, 20)
    assert [item["source"] for item in request.applied_periods(result)] == ["Calendari del centre", "Afegit manualment"]


@pytest.mark.parametrize(
    "changes",
    [
        {"weekday_minutes": {"0": 423}},
        {"weekday_minutes": {"0": 13 * 60}},
        {"target_hours": "0"},
        {"overrides": {"not-a-date": 60}},
        {"extra_periods": [{"start": "20/10/2026", "end": "19/10/2026"}]},
    ],
)
def test_request_validation_rejects_bad_input(changes):
    with pytest.raises(ValueError):
        InternshipRequest(_payload(**changes), []).calculate()


def test_target_accepts_minutes():
    request = InternshipRequest(_payload(target_hours="7", target_minutes=30), [])
    result = request.calculate()
    assert request.target_minutes == 450
    assert result.end_date == date(2026, 10, 13)
    assert result.days[-1].minutes == 30


def test_export_filename_is_ascii_slug():
    assert _export_filename({"student_name": "Núria Pérez i Puig"}) == "calendari-practiques-nuria-perez-i-puig.xlsx"
    assert _export_filename({}) == "calendari-practiques.xlsx"


def _export(result, **kwargs):
    workbook_bytes = build_internship_workbook(
        result,
        metadata=kwargs.get("metadata", {"student_name": "Núria", "company_name": "Acme"}),
        schedule=WEEKDAYS_7H,
        applied_periods=kwargs.get("applied_periods", []),
        non_working_dates=kwargs.get("non_working_dates"),
        labels_by_date=kwargs.get("labels_by_date"),
    )
    return load_workbook(workbook_bytes)


def _find(sheet, value):
    for row in sheet.iter_rows():
        for cell in row:
            cell_value = cell.value.date() if isinstance(cell.value, datetime) else cell.value
            if cell_value == value:
                return cell
    raise AssertionError(f"{value!r} not found")


def _find_day(sheet, value):
    """Find a date in the month grids (formatted as day number), not in the summary block."""

    for row in sheet.iter_rows():
        for cell in row:
            if cell.number_format == "d" and isinstance(cell.value, datetime) and cell.value.date() == value:
                return cell
    raise AssertionError(f"{value!r} not found in grid")


def test_workbook_structure_and_formulas():
    holiday = date(2026, 10, 14)
    result = build_internship_calendar(
        date(2026, 10, 12), 37 * 60 + 30, WEEKDAYS_7H, {holiday}, {date(2026, 10, 17): 150}, {holiday: "Festa local"}
    )
    workbook = _export(result)
    assert workbook.sheetnames == ["Calendari", "Detall", "Configuració"]
    grid = workbook["Calendari"]

    # Metadata and summary values.
    assert grid.cell(row=_find(grid, "Estudiant:").row, column=4).value == "Núria"
    target_cell = grid.cell(row=_find(grid, "Hores totals objectiu:").row, column=4)
    assert target_cell.value == timedelta(minutes=37 * 60 + 30)  # Time values load back as timedelta.
    assert target_cell.number_format == '[h]"h"mm'  # Totals read 4h30; day cells stay h:mm.
    planned = grid.cell(row=_find(grid, "Hores planificades:").row, column=4).value
    assert planned.startswith("=ROUND((SUM(H") and planned.endswith(")*1440,0)/1440")
    remaining = grid.cell(row=_find(grid, "Hores restants:").row, column=4).value
    assert remaining == f"=MAX(0,ROUND(($D${target_cell.row}-$D${target_cell.row + 1})*1440,0)/1440)"
    end_formula = grid.cell(row=_find(grid, "Data de fi:").row, column=4).value
    assert "_xlfn.MINIFS(Detall!" in end_formula

    # Legend lists every colour category.
    for label in ("Dia amb hores", "Sense pràctiques (0 / -)", "Sense horari"):
        _find(grid, label)

    # Each day's hours sit right below its date; non-teaching weekdays show 0 as "-".
    monday = _find_day(grid, date(2026, 10, 12))
    monday_hours = grid.cell(row=monday.row + 1, column=monday.column)
    assert monday_hours.value == timedelta(hours=7)
    assert monday_hours.number_format == '[h]:mm;-[h]:mm;"-"'
    holiday_cell = _find_day(grid, holiday)
    holiday_hours = grid.cell(row=holiday_cell.row + 1, column=holiday_cell.column)
    assert holiday_hours.value == timedelta(0)
    assert holiday_hours.number_format.endswith(';"-"')
    assert holiday_cell.comment.text.startswith("Festa local")
    saturday = _find_day(grid, date(2026, 10, 17))
    assert grid.cell(row=saturday.row + 1, column=saturday.column).value == timedelta(hours=2, minutes=30)
    assert "Hores canviades a mà" in saturday.comment.text
    # Days without a schedule are grey; the rest are coloured by conditional formatting.
    sunday = _find_day(grid, date(2026, 10, 18))
    assert grid.cell(row=sunday.row + 1, column=sunday.column).value is None
    assert sunday.fill.fgColor.rgb.endswith("E6E6E6")
    assert monday.fill.fill_type is None and saturday.fill.fill_type is None
    day_rules = {
        rule.formula[0]
        for cf in grid.conditional_formatting
        if str(cf.sqref).startswith("A") and len(cf.sqref.ranges) > 1
        for rule in cf.rules
    }
    # Rules are anchored on the first grid week (1 October) and apply across every week.
    first_week = _find_day(grid, date(2026, 10, 1)).row
    assert f"AND(ISNUMBER(A{first_week + 1}),A{first_week + 1}>0)" in day_rules
    assert f"AND(ISNUMBER(A{first_week + 1}),A{first_week + 1}=0)" in day_rules
    # Excel paints solid conditional-formatting fills with bgColor.
    assert all(
        rule.dxf.fill.bgColor.rgb.endswith(("DCEBDD", "F4D6D2")) for cf in grid.conditional_formatting for rule in cf.rules
    )

    # Hour cells only accept times between 0:00 and 12:00.
    validation = grid.data_validations.dataValidation[0]
    assert (validation.type, validation.formula1, validation.formula2) == ("time", "0", "0.5")
    assert monday_hours.coordinate in validation.sqref
    assert "4:30" in validation.error

    # Week and month totals are formulas over the hour rows.
    assert grid.cell(row=monday.row, column=8).value == f"=ROUND((SUM(A{monday.row + 1}:G{monday.row + 1}))*1440,0)/1440"
    assert f"H{monday.row}:H{monday.row + 1}" in {str(merged) for merged in grid.merged_cells.ranges}
    month_label = _find(grid, "Total octubre 2026")
    assert month_label.row > monday.row
    assert grid.cell(row=month_label.row, column=8).number_format == '[h]"h"mm'
    assert grid.cell(row=monday.row, column=8).number_format == '[h]"h"mm;-[h]"h"mm;;@'
    for label in ("Hores planificades:", "Hores restants:"):
        assert grid.cell(row=_find(grid, label).row, column=4).number_format == '[h]"h"mm'

    # Detail sheet: hours link to the grid, running and remaining totals are formulas.
    detail = workbook["Detall"]
    assert [cell.value for cell in detail[1]] == ["Data", "Dia", "Tipus", "Motiu", "Hores", "Acumulat", "Restants"]
    hours_ref = f"Calendari!{monday_hours.coordinate}"
    assert detail["E2"].value == f"=IF(ISNUMBER({hours_ref}),{hours_ref},0)"
    assert detail["F2"].value == "=ROUND((E2)*1440,0)/1440"
    assert detail["F3"].value == "=ROUND((F2+E3)*1440,0)/1440"
    assert detail["G2"].value == f"=MAX(0,ROUND((Calendari!$D${target_cell.row}-F2)*1440,0)/1440)"
    assert [detail[f"{column}2"].number_format for column in "EFG"] == ["[h]:mm", '[h]"h"mm', '[h]"h"mm']
    # The detail runs from the start date to the end of the month after the end date.
    assert detail.cell(row=2 + len(result.days), column=1).value.date() == date(2026, 10, 20)
    assert detail.cell(row=51, column=1).value.date() == date(2026, 11, 30)
    assert detail.cell(row=52, column=4).value == "Total"
    assert detail["C4"].value == "Festiu / no lectiu"
    assert detail["D4"].value == "Festa local"


def test_workbook_extends_to_following_month_for_extra_hours():
    result = _calendar(14 * 60)  # Ends on Tuesday 13/10/2026.
    local_holiday = date(2026, 11, 2)
    workbook = _export(result, non_working_dates={local_holiday}, labels_by_date={local_holiday: "Festa local"})
    grid = workbook["Calendari"]
    _find(grid, "Total novembre 2026")
    with pytest.raises(AssertionError):
        _find(grid, "Total desembre 2026")

    # Days after the end date are editable blanks; non-teaching weekdays show 0.
    after_end = _find_day(grid, date(2026, 10, 14))
    after_end_hours = grid.cell(row=after_end.row + 1, column=after_end.column)
    assert after_end_hours.value is None
    holiday = _find_day(grid, local_holiday)
    assert grid.cell(row=holiday.row + 1, column=holiday.column).value == timedelta(0)
    assert holiday.comment.text.startswith("Festa local")

    detail = workbook["Detall"]
    tail = _find(detail, local_holiday)
    assert detail.cell(row=tail.row, column=3).value == "Festiu / no lectiu"
    hours_ref = f"Calendari!{after_end_hours.coordinate}"
    assert _find(detail, f"=IF(ISNUMBER({hours_ref}),{hours_ref},0)").row == 4
    end_formula = grid.cell(row=_find(grid, "Data de fi:").row, column=4).value
    assert f"Detall!$A$2:$A${tail.row + 28}" in end_formula  # 2 Nov + 28 days = 30 Nov.


def test_workbook_never_shows_august():
    """Ending in July leaves no extra month: the grid stops on 31 July and warns about missing hours."""

    result = _calendar(200 * 60, weekday_minutes={index: 120 for index in range(5)}, start=date(2027, 6, 1))
    workbook = _export(result)
    grid = workbook["Calendari"]
    _find(grid, "Total juliol 2027")
    with pytest.raises(AssertionError):
        _find(grid, "Total agost 2027")
    remaining_row = _find(grid, "Hores restants:").row
    note = grid.cell(row=remaining_row + 2, column=1).value
    assert note.startswith(f"=IF($D${remaining_row}>0,")
    assert "31/07/2027" in note and "nou conveni" in note
    detail = workbook["Detall"]
    assert detail.cell(row=detail.max_row - 3, column=1).value.date() == date(2027, 7, 31)

    # Ending in June still gets July as the extra month, and the grid stops there.
    june = _export(_calendar(14 * 60, start=date(2027, 6, 7)))["Calendari"]
    _find(june, "Total juliol 2027")
    with pytest.raises(AssertionError):
        _find(june, "Total agost 2027")

    # Far from the deadline there is no deadline note.
    early = _export(_calendar(14 * 60))["Calendari"]
    early_remaining = _find(early, "Hores restants:").row
    assert early.cell(row=early_remaining + 2, column=1).value is None


def test_workbook_summary_instructions_and_print_area():
    workbook = _export(
        _calendar(14 * 60), metadata={"school_tutor": "Joan Tutor", "school_tutor_email": "tutor@example.com"}
    )
    grid = workbook["Calendari"]
    assert not {"school_tutor", "school_tutor_email"} & set(dict(METADATA_FIELDS))
    for value in ("Joan Tutor", "tutor@example.com", "Tutor/a del centre:"):
        with pytest.raises(AssertionError):
            _find(grid, value)

    instructions = _find(grid, "Instruccions")
    assert instructions.column > 8 and instructions.row == 1
    box = grid.cell(row=2, column=instructions.column).value
    assert "4:30" in box and "h:mm" in box and "4,5 no s'accepta" in box and "Hores restants" in box and "posteriors a la data de fi" in box

    target_row = _find(grid, "Hores totals objectiu:").row
    rules = {str(cf.sqref): cf.rules[0] for cf in grid.conditional_formatting}
    planned_rule = rules[f"D{target_row + 1}"]
    assert (planned_rule.operator, planned_rule.formula) == ("lessThan", [f"$D${target_row}"])
    remaining_rule = rules[f"D{target_row + 2}"]
    assert (remaining_rule.operator, remaining_rule.formula) == ("greaterThan", ["0"])
    assert grid.print_area.startswith("'Calendari'!$A$1:$H$")


def test_workbook_only_day_cells_and_student_data_are_editable():
    result = _calendar(14 * 60)  # 12/10/2026 to 13/10/2026; the grid runs to 30/11/2026.
    grid = _export(result)["Calendari"]
    assert grid.protection.sheet is True
    assert grid.protection.password is None

    first_day = _find_day(grid, date(2026, 10, 12))
    last_day = _find_day(grid, date(2026, 11, 30))
    before_start = _find_day(grid, date(2026, 10, 9))
    assert grid.cell(row=first_day.row + 1, column=first_day.column).protection.locked is False
    assert grid.cell(row=last_day.row + 1, column=last_day.column).protection.locked is False
    # Days before the start, padding cells outside the month and totals stay locked.
    assert grid.cell(row=before_start.row + 1, column=before_start.column).protection.locked is True
    padding = _find_day(grid, date(2026, 11, 2))  # Monday of November's first full week.
    assert grid.cell(row=padding.row - 1, column=1).protection.locked is True
    assert grid.cell(row=first_day.row, column=8).protection.locked is True
    assert grid.cell(row=_find(grid, "Estudiant:").row, column=4).protection.locked is False
    assert grid.cell(row=_find(grid, "Hores planificades:").row, column=4).protection.locked is True


def test_workbook_grid_hours_add_up_to_target_with_minutes():
    """4 h 30 min days are written as time values that sum exactly to the target."""

    result = _calendar(37 * 60 + 30, weekday_minutes={index: 270 for index in range(5)})
    grid = _export(result)["Calendari"]
    total = timedelta(0)
    for day in result.days:
        date_cell = _find_day(grid, day.date)
        value = grid.cell(row=date_cell.row + 1, column=date_cell.column).value
        total += value or timedelta(0)
    assert total == timedelta(minutes=37 * 60 + 30)


@pytest.fixture()
def client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    TestingSession = sessionmaker(bind=engine)
    with TestingSession() as db:
        year = AcademicYear(
            label="2026-27", default_start_date=date(2026, 9, 7), default_end_date=date(2027, 6, 18), is_active=True
        )
        db.add(year)
        db.flush()
        db.add(ExcludedPeriod(academic_year_id=year.id, start_date=date(2026, 10, 13), end_date=date(2026, 10, 13), label="Festa"))
        db.commit()

    def override_db():
        with TestingSession() as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    yield TestClient(app), TestingSession
    app.dependency_overrides.pop(get_db, None)


def test_page_renders_with_nav_and_school_periods(client):
    test_client, _ = client
    response = test_client.get("/practiques")
    assert response.status_code == 200
    assert 'href="/practiques" class="site-nav-link" aria-current="page"' in response.text
    assert "Festa" in response.text
    assert "Planificador d&#39;hores de pràctiques" in response.text
    assert "Correu del tutor/a del centre" not in response.text
    assert "Tutor/a del centre" not in response.text
    assert "Pots canviar les hores de qualsevol dia" in response.text
    assert "<h2>Dies sense pràctiques</h2>" in response.text
    assert "Estudiant" in response.text and "Alumne/a" not in response.text
    assert '<div class="restore-notice" id="restore-notice" role="status" hidden>' in response.text
    home = test_client.get("/")
    assert 'href="/" class="site-nav-link" aria-current="page"' in home.text
    assert "Planificador curricular" in home.text


def test_calcul_endpoint_returns_days_and_end_date(client):
    test_client, Session = client
    with Session() as db:
        period_id = db.query(ExcludedPeriod).one().id
    response = test_client.post("/practiques/calcul", json=_payload(excluded_period_ids=[period_id]))
    data = response.json()
    assert data["ok"] is True
    assert data["end_date"] == "2026-10-19"
    assert data["days"][1]["kind"] == "holiday"
    assert data["warning"] is None


def test_calcul_endpoint_reports_validation_errors_in_catalan(client):
    test_client, _ = client
    response = test_client.post("/practiques/calcul", json=_payload(start_date="31/02/2026"))
    assert response.status_code == 400
    assert response.json() == {"ok": False, "error": "Introdueix una data en format D/M/AA o D/M/AAAA."}
    response = test_client.post("/practiques/calcul", content=b"not json", headers={"Content-Type": "application/json"})
    assert response.json()["error"] == "Dades de càlcul no vàlides."


def test_calcul_reports_hours_that_do_not_fit_before_august(client):
    test_client, _ = client
    response = test_client.post(
        "/practiques/calcul", json=_payload(start_date="01/07/2027", target_hours="350", weekday_minutes={"0": 420})
    )
    data = response.json()
    assert data["complete"] is False
    assert data["end_date"] == "2027-07-31"
    assert data["deadline_display"] == "31/07/2027"
    assert data["last_working_display"] == "26/07/2027"  # Last Monday of July.
    assert data["pending_display"] == "322 h"  # Four Mondays of 7 h (28 h) fit.
    assert test_client.post("/practiques/calcul", json=_payload()).json()["complete"] is True
    august = test_client.post("/practiques/calcul", json=_payload(start_date="03/08/2027"))
    assert august.status_code == 400 and "agost" in august.json()["error"]


def test_calcul_warns_when_running_past_configured_years(client):
    test_client, _ = client
    response = test_client.post("/practiques/calcul", json=_payload(start_date="01/06/2027", target_hours="350"))
    assert "18/06/2027" in response.json()["warning"]


def test_export_endpoint_streams_workbook(client):
    test_client, _ = client
    response = test_client.post("/practiques/export", data={"payload": json.dumps(_payload())})
    assert response.status_code == 200
    assert 'filename="calendari-practiques-nuria-perez.xlsx"' in response.headers["content-disposition"]


def test_export_error_rerenders_page_without_script_injection(client):
    test_client, _ = client
    payload = _payload(target_hours="0", metadata={"student_name": "</script><script>alert(1)</script>"})
    response = test_client.post("/practiques/export", data={"payload": json.dumps(payload)})
    assert response.status_code == 400
    assert "Indica el total" in response.text
    assert "<script>alert(1)</script>" not in response.text


def test_schedule_change_applies_from_its_date():
    """7 h/day the first week, then 4 h on Monday to Thursday only."""

    change = date(2026, 10, 19)
    schedule = InternshipSchedule(WEEKDAYS_7H, [(change, {0: 240, 1: 240, 2: 240, 3: 240})])
    result = _calendar(35 * 60 + 16 * 60, weekday_minutes=schedule)
    minutes = {day.date: day.minutes for day in result.days}
    assert minutes[date(2026, 10, 16)] == 420
    assert minutes[change] == 240
    assert result.end_date == date(2026, 10, 22)
    assert _calendar(35 * 60 + 17 * 60, weekday_minutes=schedule).end_date == date(2026, 10, 26)  # No Friday hours.
    assert {day.schedule_index for day in result.days if day.date < change} == {0}
    assert {day.schedule_index for day in result.days if day.date >= change} == {1}


def test_schedule_changes_are_applied_in_date_order():
    later, earlier = date(2026, 11, 2), date(2026, 10, 26)
    schedule = InternshipSchedule(WEEKDAYS_7H, [(later, {0: 60}), (earlier, {0: 120})])
    assert schedule.minutes_on(date(2026, 10, 19)) == 420
    assert schedule.minutes_on(earlier) == 120
    assert schedule.minutes_on(later) == 60


def test_override_on_a_later_schedule_keeps_that_schedule_as_default():
    change = date(2026, 10, 19)
    schedule = InternshipSchedule(WEEKDAYS_7H, [(change, {index: 240 for index in range(5)})])
    result = _calendar(100 * 60, weekday_minutes=schedule, overrides={date(2026, 10, 20): 60})
    day = next(day for day in result.days if day.date == date(2026, 10, 20))
    assert (day.kind, day.default_minutes, day.minutes) == (DAY_KIND_OVERRIDE, 240, 60)


def test_request_parses_schedule_changes():
    request = InternshipRequest(
        _payload(target_hours="51", schedule_changes=[{"start": "19/10/2026", "weekday_minutes": {"0": 240, "1": 240, "2": 240, "3": 240}}]),
        [],
    )
    result = request.calculate()
    assert result.end_date == date(2026, 10, 22)


@pytest.mark.parametrize(
    "changes",
    [
        [{"start": "12/10/2026", "weekday_minutes": {"0": 60}}],
        [{"start": "01/10/2026", "weekday_minutes": {"0": 60}}],
        [{"start": "19/10/2026", "weekday_minutes": {"0": 60}}, {"start": "19/10/2026", "weekday_minutes": {"0": 120}}],
        [{"start": "19/10/2026", "weekday_minutes": {"0": 13 * 60}}],
        [{"start": "", "weekday_minutes": {"0": 60}}],
    ],
)
def test_request_rejects_bad_schedule_changes(changes):
    with pytest.raises(ValueError):
        InternshipRequest(_payload(schedule_changes=changes), []).calculate()


def test_calcul_endpoint_reports_schedule_of_each_day(client):
    test_client, _ = client
    response = test_client.post(
        "/practiques/calcul",
        json=_payload(target_hours="50", schedule_changes=[{"start": "19/10/2026", "weekday_minutes": {"0": 300}}]),
    )
    days = {day["date"]: day for day in response.json()["days"]}
    assert days["2026-10-16"]["schedule"] == 0
    assert days["2026-10-19"]["schedule"] == 1
    assert days["2026-10-19"]["minutes"] == 300


def test_workbook_lists_every_schedule_with_its_dates():
    change = date(2026, 10, 19)
    schedule = InternshipSchedule(WEEKDAYS_7H, [(change, {0: 240})])
    result = _calendar(60 * 60, weekday_minutes=schedule)
    workbook = load_workbook(
        build_internship_workbook(result, metadata={}, schedule=schedule, applied_periods=[])
    )
    config = workbook["Configuració"]
    first = _find(config, "Del 12/10/2026 al 18/10/2026")
    second = _find(config, "A partir del 19/10/2026")
    assert config.cell(row=first.row + 1, column=2).value == "7 h"
    assert config.cell(row=second.row + 1, column=2).value == "4 h"
    assert config.cell(row=second.row + 2, column=2).value == "-"
    # Tuesday after the change has no schedule, so its grid cell stays blank.
    assert _find_day(workbook["Calendari"], date(2026, 10, 20)).offset(row=1).value is None


def _html_export(test_client, payload):
    return test_client.post("/practiques/export-html", data={"payload": json.dumps(payload)})


def test_html_export_downloads_a_self_contained_read_only_copy(client):
    test_client, Session = client
    with Session() as db:
        period_id = db.query(ExcludedPeriod).one().id
    response = _html_export(
        test_client, _payload(excluded_period_ids=[period_id], overrides={"2026-10-17": 240}, metadata={"student_name": "Núria Pérez", "company_name": "Acme"})
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert 'filename="calendari-practiques-nuria-perez.html"' in response.headers["content-disposition"]
    html = response.text
    # Nothing is loaded from the server and nothing runs: the only script block is inert JSON.
    assert "../fonts/" not in html and "data:font/woff2;base64," in html
    assert "/static/" not in html
    assert html.count("<script") == 1 and '<script type="application/json"' in html
    assert "Còpia de només lectura" in html
    assert "<dd>Núria Pérez</dd>" in html and "<dd>Acme</dd>" in html
    assert "Octubre 2026" in html
    # 12/10 Monday start, 13/10 school holiday, Saturday 17/10 given 4 h by hand: 35 h end on 19/10.
    assert 'class="month-day day-working" data-date="2026-10-12"' in html
    assert 'class="month-day day-holiday" data-date="2026-10-13"' in html
    assert 'class="month-day day-override" data-date="2026-10-17"' in html
    assert 'class="month-day day-weekend" data-date="2026-10-18"' in html
    assert 'class="month-day day-working day-end" data-date="2026-10-19"' in html
    assert "dt. 13/10/2026 · Festiu / no lectiu · Festa · sense hores" in html
    plan = json.loads(html.split('id="tempo-internship-plan">')[1].split("</script>")[0])
    assert plan["format"] == "tempo.internship" and plan["plan"]["overrides"] == {"2026-10-17": 240}


def test_html_export_escapes_user_text(client):
    test_client, _ = client
    name = "</script><script>alert(1)</script>"
    response = _html_export(test_client, _payload(metadata={"student_name": name}, extra_periods=[{"start": "14/10/2026", "label": "<b>x</b>"}]))
    assert response.status_code == 200
    assert "<script>alert(1)</script>" not in response.text
    assert "<b>x</b>" not in response.text


def test_html_export_error_rerenders_planner(client):
    test_client, _ = client
    response = _html_export(test_client, _payload(target_hours="0"))
    assert response.status_code == 400
    assert "Indica el total" in response.text
    assert test_client.get("/practiques/export-html", follow_redirects=False).status_code == 303


def test_html_export_and_planner_share_day_labels_and_legend(client):
    test_client, _ = client
    page = test_client.get("/practiques").text
    exported = _html_export(test_client, _payload()).text
    assert json.loads(page.split("const KIND_LABELS = ")[1].split(";\n")[0]) == CALENDAR_KIND_LABELS
    for kind, label in CALENDAR_KIND_LABELS.items():
        swatch = f'<span class="legend-swatch day-{kind}"></span>{label}'
        assert swatch in page and swatch in exported
    assert kind_classes_in_css() >= {f"day-{kind}" for kind in CALENDAR_KIND_LABELS}


def kind_classes_in_css():
    css = inline_stylesheet()
    return {kind for kind in (f"day-{name}" for name in CALENDAR_KIND_LABELS) if f".{kind} {{" in css}


def test_month_views_mark_schedule_changes_and_partial_last_day():
    schedule = InternshipSchedule(WEEKDAYS_7H, [(date(2026, 10, 19), {0: 240, 1: 240, 2: 240, 3: 240, 4: 240})])
    result = _calendar(37 * 60 + 30, weekday_minutes=schedule)
    cells = {cell["day"].date: cell for month in build_month_views(result) for cell in month["cells"] if cell["day"]}
    assert "day-schedule-start" in cells[date(2026, 10, 19)]["classes"]
    assert "comença el període 2" in cells[date(2026, 10, 19)]["title"]
    assert "període 1" in cells[date(2026, 10, 12)]["title"]
    last = cells[result.end_date]
    assert last["hours"] == "2h30" and "darrer dia parcial" in last["title"] and "day-end" in last["classes"]
    assert build_month_views(result)[0]["leading_blanks"] == 3  # 1 October 2026 is a Thursday.


def test_html_export_groups_header_details_in_columns(client):
    test_client, _ = client
    metadata = {
        "student_name": "Núria",
        "group_name": "GAJ2",
        "company_name": "Acme",
        "company_tutor": "Pau",
        "company_tutor_email": "pau@acme.cat",
    }
    html = _html_export(test_client, _payload(metadata=metadata)).text
    columns = [
        [dt.split("</dt>")[0] for dt in column.split("<dt>")[1:]]
        for column in html.split('<dl class="snapshot-column">')[1:]
    ]
    columns[-1] = columns[-1][:2]
    assert columns == [
        ["Estudiant", "Grup"],
        ["Empresa"],
        ["Tutor/a de l&#39;empresa", "Correu del tutor/a de l&#39;empresa"],
        ["Data d&#39;inici", "Hores totals"],
    ]
    # With no student or company details, only the dates remain.
    assert _html_export(test_client, _payload(metadata={})).text.count('<dl class="snapshot-column">') == 1
