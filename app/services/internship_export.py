"""XLSX export for internship calendars.

The workbook has three sheets:

- "Calendari": one Monday-to-Sunday grid per month, from the start month to the
  month after the end date so hours can be added past the planned end. Each
  week uses two rows, the dates and the editable hours below them. Week, month
  and overall totals are formulas, so editing any day's hours recalculates
  everything, and conditional formatting colours each day by its hours.
  The grid never goes past 31 July: internships cannot run in August.
  Hours are spreadsheet time values shown as h:mm, so a day is typed as 4:30.
- "Detall": one row per editable date. Hours link back to the grid and the
  running and remaining totals are formulas.
- "Configuració": the weekly schedules and the non-teaching periods that were applied.
"""

from __future__ import annotations

import calendar as month_calendar
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from app.date_utils import format_display_date, format_month_label
from app.services.hours import format_minutes_label
from app.services.internship import (
    DAY_KIND_HOLIDAY,
    DAY_KIND_OVERRIDE,
    DAY_KIND_WEEKEND,
    DAY_KIND_WORKING,
    MAX_DAILY_HOURS,
    InternshipCalendar,
    InternshipSchedule,
    classify_day,
    default_minutes_for,
)


KIND_FILLS = {
    DAY_KIND_WORKING: "DCEBDD",
    DAY_KIND_HOLIDAY: "F4D6D2",
    DAY_KIND_WEEKEND: "E6E6E6",
    DAY_KIND_OVERRIDE: "F6E2B3",
}
KIND_LABELS = {
    DAY_KIND_WORKING: "Laborable",
    DAY_KIND_HOLIDAY: "Festiu / no lectiu",
    DAY_KIND_WEEKEND: "Sense horari (cap de setmana)",
    DAY_KIND_OVERRIDE: "Hores canviades a mà",
}
# The grid itself only distinguishes days with hours (green), days without
# practice time (0 shown as "-", red) and days without a schedule (grey, blank).
HOURS_FILL_COLOR = KIND_FILLS[DAY_KIND_WORKING]
NO_HOURS_FILL_COLOR = KIND_FILLS[DAY_KIND_HOLIDAY]
NO_SCHEDULE_FILL_COLOR = KIND_FILLS[DAY_KIND_WEEKEND]
GRID_LEGEND = (
    ("Dia amb hores", HOURS_FILL_COLOR),
    ("Sense pràctiques (0 / -)", NO_HOURS_FILL_COLOR),
    ("Sense horari", NO_SCHEDULE_FILL_COLOR),
)
OUTSIDE_FONT_COLOR = "A0A69A"
MONTH_TITLE_FILL = PatternFill(fill_type="solid", fgColor="1E3A2D")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="F2ECE0")
THIN_SIDE = Side(style="thin", color="B7BDA8")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
END_DAY_SIDE = Side(style="medium", color="1E3A2D")
WEEKDAY_SHORT = ("dl.", "dt.", "dc.", "dj.", "dv.", "ds.", "dg.")
GRID_COLUMNS = 8  # Monday..Sunday plus the weekly total.
GRID_COLUMN_WIDTH = 11
METADATA_FIELDS = (
    ("student_name", "Estudiant"),
    ("training_cycle", "Cicle formatiu"),
    ("group_name", "Grup"),
    ("company_name", "Empresa"),
    ("company_tutor", "Tutor/a de l'empresa"),
    ("company_tutor_email", "Correu del tutor/a de l'empresa"),
)
GRID_INSTRUCTIONS = (
    "Podeu canviar les hores de qualsevol dia escrivint-les a la casella de sota la data. "
    "Els totals, les hores restants i la data de fi es recalculen automàticament.",
    "També podeu afegir hores en dies posteriors a la data de fi prevista: el calendari inclou el mes següent, "
    "fins al 31 de juliol (a l'agost no es poden fer pràctiques).",
    "Escriviu les hores en format hores:minuts (h:mm): 4:30 per a 4 h 30 min, 7:00 per a 7 h o 0:45 per a 45 min. "
    "No feu servir decimals: 4,5 no s'accepta. Cada dia admet entre 0:00 i 12:00.",
    "Per indicar que un dia no hi ha pràctiques, escriviu 0: es mostra com «-». Per treure un dia, esborreu-ne el contingut.",
    "Els dies amb hores es mostren en verd i els dies sense pràctiques (0, que es mostra «-»), en vermell.",
    "Els totals es calculen sols i es mostren com 4h30 (p. ex. 350h00 són 350 hores).",
    "El full està protegit perquè només es puguin editar les hores dels dies i les dades de l'estudiant i de l'empresa. "
    "Per desprotegir-lo: Revisió › Desprotegeix el full (no té contrasenya).",
    "«Hores planificades» es marca en vermell si no arriba a les hores totals objectiu, "
    "i «Hores restants» si encara queden hores per fer.",
)
UNLOCKED = Protection(locked=False)
# The instructions sit to the right of the printable grid, after a spacer column.
INSTRUCTIONS_FIRST_COLUMN = GRID_COLUMNS + 2
INSTRUCTIONS_LAST_COLUMN = INSTRUCTIONS_FIRST_COLUMN + 7


# Hours are stored as spreadsheet time values (fractions of a day) so day cells
# accept and show "4:30"; calculated totals read "4h30" so they are not mistaken
# for editable cells. [h] keeps totals above 24 h from wrapping around.
MINUTES_PER_DAY = 24 * 60
TIME_FORMAT = "[h]:mm"
DAY_TIME_FORMAT = '[h]:mm;-[h]:mm;"-"'
TOTAL_FORMAT = '[h]"h"mm'
WEEK_TOTAL_FORMAT = '[h]"h"mm;-[h]"h"mm;;@'  # Hides zero week totals outside the internship.


def minutes_to_time(minutes: int) -> float:
    """Return minutes as a spreadsheet time value (a fraction of a day)."""

    return minutes / MINUTES_PER_DAY


def _round_to_minute(expression: str) -> str:
    """Wrap a time formula so sums of time values compare exactly against the target."""

    return f"ROUND(({expression})*{MINUTES_PER_DAY},0)/{MINUTES_PER_DAY}"


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _rule_fill(color: str) -> PatternFill:
    """Conditional-formatting fill: Excel paints solid rule fills with bgColor, not fgColor."""

    return PatternFill(fill_type="solid", fgColor=color, bgColor=color)


def _next_month_start(current: date) -> date:
    return (current.replace(day=28) + timedelta(days=4)).replace(day=1)


def _month_starts(start: date, end: date) -> list[date]:
    months: list[date] = []
    current = start.replace(day=1)
    while current <= end:
        months.append(current)
        current = _next_month_start(current)
    return months


def calendar_grid_end(result: InternshipCalendar) -> date:
    """Last date in the exported grid: the end of the month after the end date, but never past 31 July."""

    return min(_next_month_start(_next_month_start(result.end_date)) - timedelta(days=1), result.deadline)


def _label_value_row(sheet, row: int, label: str, value, *, bold_value: bool = False) -> None:
    """Write a 'label: value' line spanning A:C and D:H of the grid sheet."""

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=GRID_COLUMNS)
    label_cell = sheet.cell(row=row, column=1, value=f"{label}:")
    label_cell.font = Font(bold=True)
    value_cell = sheet.cell(row=row, column=4, value=value)
    value_cell.alignment = Alignment(horizontal="left")
    if bold_value:
        value_cell.font = Font(bold=True)


def _grid_day_value(current: date, minutes: int | None, schedule: InternshipSchedule, non_working_dates: set[date]) -> int | None:
    """Return the grid value of a day in whole minutes, or None for a blank cell.

    Non-teaching days that would normally have practice time show 0 (rendered
    as "-") so they stand out; days without a schedule stay blank.
    """

    if minutes:
        return minutes
    # `minutes == 0` on a scheduled day not marked non-teaching means it was set to zero by hand.
    if schedule.minutes_on(current) and (current in non_working_dates or minutes == 0):
        return 0
    return None


def _add_day_colour_rules(sheet, date_ranges: list[str], hour_ranges: list[str]) -> None:
    """Colour each day green when it has hours and red when it shows 0 ("-").

    The date cells check the hours cell right below them, so both cells of a
    day change colour together as soon as its hours are edited.
    """

    for ranges, offset in ((date_ranges, 1), (hour_ranges, 0)):
        if not ranges:
            continue
        first = ranges[0].split(":")[0]
        column = "".join(char for char in first if char.isalpha())
        target = f"{column}{int(first[len(column):]) + offset}"
        sqref = " ".join(ranges)
        sheet.conditional_formatting.add(
            sqref, FormulaRule(formula=[f"AND(ISNUMBER({target}),{target}>0)"], fill=_rule_fill(HOURS_FILL_COLOR))
        )
        sheet.conditional_formatting.add(
            sqref, FormulaRule(formula=[f"AND(ISNUMBER({target}),{target}=0)"], fill=_rule_fill(NO_HOURS_FILL_COLOR))
        )


def _add_hours_validation(sheet, cells: list[str]) -> None:
    """Only accept times between 0:00 and the daily maximum, so "4" or "4,5" are rejected."""

    validation = DataValidation(
        type="time",
        operator="between",
        formula1="0",
        formula2=str(MAX_DAILY_HOURS / 24),
        allow_blank=True,
        showErrorMessage=True,
        errorStyle="stop",
        errorTitle="Format d'hores",
        error=f"Escriviu les hores en format h:mm, p. ex. 4:30 o 7:00 (entre 0:00 i {MAX_DAILY_HOURS}:00).",
    )
    sheet.add_data_validation(validation)
    for cell in cells:
        validation.add(cell)


def _schedule_sections(result: InternshipCalendar, schedule: InternshipSchedule) -> list[tuple[str, dict[int, int]]]:
    """Title each weekly schedule with the dates it covers; a single schedule needs no title."""

    if not schedule.changes:
        return [("", schedule.base)]
    starts = [result.start_date] + [start for start, _ in schedule.changes]
    sections = []
    for index, weekday_minutes in enumerate(schedule.weekly_minutes):
        start = format_display_date(starts[index])
        if index + 1 < len(starts):
            title = f"Del {start} al {format_display_date(starts[index + 1] - timedelta(days=1))}"
        else:
            title = f"A partir del {start}"
        sections.append((title, weekday_minutes))
    return sections


def build_internship_workbook(
    result: InternshipCalendar,
    *,
    metadata: dict[str, str],
    schedule: InternshipSchedule | dict[int, int],
    applied_periods: list[dict[str, str]],
    non_working_dates: set[date] | None = None,
    labels_by_date: dict[date, str] | None = None,
) -> BytesIO:
    """Create the internship workbook and return it as an in-memory file.

    `applied_periods` lists the non-working periods that were applied, as
    dictionaries with `start`, `end`, `label` and `source` display strings.
    `non_working_dates` and `labels_by_date` mark the non-teaching days after
    the planned end date, which the grid also shows.
    """

    schedule = InternshipSchedule.coerce(schedule)
    non_working_dates = non_working_dates or set()
    labels_by_date = labels_by_date or {}

    workbook = Workbook()
    grid = workbook.active
    grid.title = "Calendari"
    detail = workbook.create_sheet("Detall")
    config = workbook.create_sheet("Configuració")

    for column in range(1, GRID_COLUMNS + 1):
        grid.column_dimensions[get_column_letter(column)].width = GRID_COLUMN_WIDTH
    grid.column_dimensions[get_column_letter(GRID_COLUMNS + 1)].width = 3
    for column in range(INSTRUCTIONS_FIRST_COLUMN, INSTRUCTIONS_LAST_COLUMN + 1):
        grid.column_dimensions[get_column_letter(column)].width = GRID_COLUMN_WIDTH

    grid.merge_cells(start_row=1, start_column=1, end_row=1, end_column=GRID_COLUMNS)
    title_cell = grid.cell(row=1, column=1, value="Calendari de pràctiques")
    title_cell.font = Font(bold=True, size=14)
    grid.row_dimensions[1].height = 22

    row = 3
    for key, label in METADATA_FIELDS:
        _label_value_row(grid, row, label, (metadata.get(key) or "").strip() or None)
        grid.cell(row=row, column=4).protection = UNLOCKED
        row += 1

    row += 1
    _label_value_row(grid, row, "Data d'inici", result.start_date)
    grid.cell(row=row, column=4).number_format = "DD/MM/YYYY"
    row += 1
    target_row = row
    _label_value_row(grid, target_row, "Hores totals objectiu", minutes_to_time(result.target_minutes), bold_value=True)
    grid.cell(row=target_row, column=4).number_format = TOTAL_FORMAT
    row += 1
    planned_row = row
    _label_value_row(grid, planned_row, "Hores planificades", None, bold_value=True)
    row += 1
    remaining_row = row
    _label_value_row(grid, remaining_row, "Hores restants", None, bold_value=True)
    row += 1
    end_row = row
    _label_value_row(grid, end_row, "Data de fi", None, bold_value=True)
    row += 1
    if calendar_grid_end(result) == result.deadline:
        # The grid stops at 31 July, so hours still missing there need a new agreement.
        # A formula keeps the notice in step with edits to the grid.
        grid.merge_cells(start_row=row, start_column=1, end_row=row, end_column=GRID_COLUMNS)
        deadline_note = grid.cell(
            row=row,
            column=1,
            value=(
                f'=IF($D${remaining_row}>0,"Amb les hores actuals no s\'arriba a l\'objectiu abans del '
                f"{format_display_date(result.deadline)} (a l'agost no es poden fer pràctiques). "
                f'Les hores restants s\'hauran de fer amb un nou conveni el curs vinent.","")'
            ),
        )
        deadline_note.font = Font(bold=True, color="B3473A")
        deadline_note.alignment = Alignment(wrap_text=True, vertical="center")
        grid.row_dimensions[row].height = 30
    row += 1

    target_ref = f"$D${target_row}"
    planned_ref = f"$D${planned_row}"
    remaining_ref = f"$D${remaining_row}"
    grid.conditional_formatting.add(
        planned_ref.replace("$", ""), CellIsRule(operator="lessThan", formula=[target_ref], fill=_rule_fill(NO_HOURS_FILL_COLOR))
    )
    grid.conditional_formatting.add(
        remaining_ref.replace("$", ""), CellIsRule(operator="greaterThan", formula=["0"], fill=_rule_fill(NO_HOURS_FILL_COLOR))
    )

    # Instructions box in the top right corner, next to the summary block.
    instructions_title = grid.cell(row=1, column=INSTRUCTIONS_FIRST_COLUMN, value="Instruccions")
    instructions_title.font = Font(bold=True, size=12)
    grid.merge_cells(
        start_row=2, start_column=INSTRUCTIONS_FIRST_COLUMN, end_row=end_row + 2, end_column=INSTRUCTIONS_LAST_COLUMN
    )
    instructions = grid.cell(row=2, column=INSTRUCTIONS_FIRST_COLUMN, value="\n".join(f"• {line}" for line in GRID_INSTRUCTIONS))
    instructions.alignment = Alignment(wrap_text=True, vertical="top")
    instructions.font = Font(size=10)
    instructions.fill = HEADER_FILL

    legend_row = row
    grid.cell(row=legend_row, column=1, value="Llegenda:").font = Font(bold=True)
    legend_column = 2
    for label, color in GRID_LEGEND:
        grid.merge_cells(start_row=legend_row, start_column=legend_column, end_row=legend_row, end_column=legend_column + 1)
        cell = grid.cell(row=legend_row, column=legend_column, value=label)
        cell.fill = _fill(color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.font = Font(size=9)
        legend_column += 2
    grid.row_dimensions[legend_row].height = 26
    row += 2

    days_by_date = {day.date: day for day in result.days}
    grid_end = calendar_grid_end(result)
    hour_cell_by_date: dict[date, str] = {}
    month_total_cells: list[str] = []
    date_ranges: list[str] = []
    hour_ranges: list[str] = []

    for month_start in _month_starts(result.start_date, grid_end):
        grid.merge_cells(start_row=row, start_column=1, end_row=row, end_column=GRID_COLUMNS)
        month_cell = grid.cell(row=row, column=1, value=format_month_label(month_start))
        month_cell.font = Font(bold=True, color="FFFFFF")
        month_cell.fill = MONTH_TITLE_FILL
        month_cell.alignment = Alignment(horizontal="center", vertical="center")
        row += 1

        for column, weekday in enumerate(WEEKDAY_SHORT, start=1):
            header = grid.cell(row=row, column=column, value=weekday)
            header.font = Font(bold=True)
            header.fill = HEADER_FILL
            header.alignment = Alignment(horizontal="center")
            header.border = THIN_BORDER
        total_header = grid.cell(row=row, column=GRID_COLUMNS, value="Setmana")
        total_header.font = Font(bold=True)
        total_header.fill = HEADER_FILL
        total_header.alignment = Alignment(horizontal="center")
        total_header.border = THIN_BORDER
        row += 1

        week_total_cells: list[str] = []
        for week in month_calendar.Calendar(firstweekday=0).monthdatescalendar(month_start.year, month_start.month):
            date_row, hours_row = row, row + 1
            date_ranges.append(f"A{date_row}:G{date_row}")
            hour_ranges.append(f"A{hours_row}:G{hours_row}")
            for column, current in enumerate(week, start=1):
                if current.month != month_start.month:
                    continue
                date_cell = grid.cell(row=date_row, column=column, value=current)
                date_cell.number_format = "d"
                date_cell.alignment = Alignment(horizontal="right", vertical="top")
                date_cell.font = Font(size=8, color="607165")
                hours_cell = grid.cell(row=hours_row, column=column)
                hours_cell.alignment = Alignment(horizontal="center", vertical="center")
                hours_cell.font = Font(bold=True)
                date_cell.border = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE)
                hours_cell.border = Border(left=THIN_SIDE, right=THIN_SIDE, bottom=THIN_SIDE)

                if current < result.start_date:
                    date_cell.font = Font(size=8, color=OUTSIDE_FONT_COLOR)
                    continue
                day = days_by_date.get(current)
                minutes = day.minutes if day is not None else None
                value = _grid_day_value(current, minutes, schedule, non_working_dates)
                if value is not None:
                    hours_cell.value = minutes_to_time(value)
                hours_cell.number_format = DAY_TIME_FORMAT
                if value is None and not schedule.minutes_on(current):
                    # Conditional formatting colours green/red once hours are typed in.
                    date_cell.fill = _fill(NO_SCHEDULE_FILL_COLOR)
                    hours_cell.fill = _fill(NO_SCHEDULE_FILL_COLOR)

                notes = []
                label = day.label if day is not None else labels_by_date.get(current, "")
                if label:
                    notes.append(label)
                if day is not None and day.kind == DAY_KIND_OVERRIDE:
                    notes.append(f"Hores canviades a mà (horari habitual: {format_minutes_label(day.default_minutes)})")
                if notes:
                    date_cell.comment = Comment("\n".join(notes), "Tempo")
                if current == result.end_date:
                    date_cell.border = Border(left=END_DAY_SIDE, right=END_DAY_SIDE, top=END_DAY_SIDE)
                    hours_cell.border = Border(left=END_DAY_SIDE, right=END_DAY_SIDE, bottom=END_DAY_SIDE)
                hour_cell_by_date[current] = f"{get_column_letter(column)}{hours_row}"

            # One tall weekly-total cell spanning the date and hours rows.
            grid.merge_cells(start_row=date_row, start_column=GRID_COLUMNS, end_row=hours_row, end_column=GRID_COLUMNS)
            week_total = grid.cell(row=date_row, column=GRID_COLUMNS, value=f"={_round_to_minute(f'SUM(A{hours_row}:G{hours_row})')}")
            week_total.alignment = Alignment(horizontal="center", vertical="center")
            # Weeks outside the internship would otherwise show a column of zeros.
            week_total.number_format = WEEK_TOTAL_FORMAT
            week_total.font = Font(bold=True)
            week_total.border = THIN_BORDER
            grid.cell(row=hours_row, column=GRID_COLUMNS).border = THIN_BORDER
            week_total_cells.append(f"H{date_row}")
            grid.row_dimensions[date_row].height = 12
            grid.row_dimensions[hours_row].height = 20
            row += 2

        grid.merge_cells(start_row=row, start_column=1, end_row=row, end_column=GRID_COLUMNS - 1)
        month_label = grid.cell(row=row, column=1, value=f"Total {format_month_label(month_start).lower()}")
        month_label.font = Font(bold=True)
        month_label.alignment = Alignment(horizontal="right")
        week_sum = f"SUM({','.join(week_total_cells)})"
        month_total = grid.cell(row=row, column=GRID_COLUMNS, value=f"={_round_to_minute(week_sum)}")
        month_total.number_format = TOTAL_FORMAT
        month_total.font = Font(bold=True)
        month_total.alignment = Alignment(horizontal="center")
        month_total.border = THIN_BORDER
        month_total_cells.append(f"H{row}")
        row += 2

    _add_day_colour_rules(grid, date_ranges, hour_ranges)
    _add_hours_validation(grid, list(hour_cell_by_date.values()))
    for coordinate in hour_cell_by_date.values():
        grid[coordinate].protection = UNLOCKED
    # Lock everything else (padding cells, totals, formulas). No password: it only
    # prevents accidental edits, and resizing rows and columns stays allowed.
    grid.protection.sheet = True
    grid.protection.formatColumns = False
    grid.protection.formatRows = False
    month_sum = f"SUM({','.join(month_total_cells)})"
    grid.cell(row=planned_row, column=4).value = f"={_round_to_minute(month_sum)}"
    # Spreadsheets show negative times as "####", so extra hours leave 0:00 remaining.
    grid.cell(row=remaining_row, column=4).value = f"=MAX(0,{_round_to_minute(f'{target_ref}-{planned_ref}')})"
    for cell_row in (target_row, planned_row, remaining_row):
        grid.cell(row=cell_row, column=4).alignment = Alignment(horizontal="left")
        grid.cell(row=cell_row, column=4).number_format = TOTAL_FORMAT

    # Detail sheet: one row per editable date, linked to the grid's hour cells.
    detail_headers = ("Data", "Dia", "Tipus", "Motiu", "Hores", "Acumulat", "Restants")
    detail.append(detail_headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
    detail.freeze_panes = "A2"
    first_detail_row = 2
    detail_dates = sorted(hour_cell_by_date)
    for index, current in enumerate(detail_dates):
        detail_row = first_detail_row + index
        day = days_by_date.get(current)
        if day is not None:
            kind, label = day.kind, day.label
        else:
            default_minutes = default_minutes_for(current, schedule, non_working_dates)
            kind = classify_day(current, default_minutes, default_minutes, False, non_working_dates)
            label = labels_by_date.get(current, "")
        hours_ref = f"Calendari!{hour_cell_by_date[current]}"
        previous = f"F{detail_row - 1}+" if index else ""
        detail.append(
            [
                current,
                WEEKDAY_SHORT[current.weekday()],
                KIND_LABELS[kind],
                label or None,
                f"=IF(ISNUMBER({hours_ref}),{hours_ref},0)",
                f"={_round_to_minute(f'{previous}E{detail_row}')}",
                f"=MAX(0,{_round_to_minute(f'Calendari!{target_ref}-F{detail_row}')})",
            ]
        )
        # Dates after the planned end stay unfilled: they only count if hours are added.
        fill = _fill(KIND_FILLS[kind]) if day is not None else None
        for column in range(1, len(detail_headers) + 1):
            cell = detail.cell(row=detail_row, column=column)
            if fill is not None:
                cell.fill = fill
            cell.border = THIN_BORDER
            if column != 4:
                cell.alignment = Alignment(horizontal="center")
        detail.cell(row=detail_row, column=1).number_format = "DD/MM/YYYY"
        # Daily hours match the grid; running and remaining totals use the total format.
        detail.cell(row=detail_row, column=5).number_format = TIME_FORMAT
        for column in (6, 7):
            detail.cell(row=detail_row, column=column).number_format = TOTAL_FORMAT
    last_detail_row = first_detail_row + len(detail_dates) - 1
    total_row = last_detail_row + 1
    detail.cell(row=total_row, column=4, value="Total").font = Font(bold=True)
    detail_total = detail.cell(row=total_row, column=5, value=f"={_round_to_minute(f'SUM(E{first_detail_row}:E{last_detail_row})')}")
    detail_total.number_format = TOTAL_FORMAT
    detail_total.font = Font(bold=True)
    detail_total.alignment = Alignment(horizontal="center")
    detail.cell(row=total_row + 2, column=1, value="Les hores d'aquest full provenen del full Calendari: modifiqueu-les allà.").font = Font(
        italic=True, color="607165"
    )
    for column, width in zip("ABCDEFG", (12, 6, 26, 32, 9, 10, 10)):
        detail.column_dimensions[column].width = width

    # The end date is the first listed date whose running total reaches the
    # target. MINIFS keeps it live when hours are edited in the grid.
    dates_range = f"Detall!$A${first_detail_row}:$A${last_detail_row}"
    cumulative_range = f"Detall!$F${first_detail_row}:$F${last_detail_row}"
    end_cell = grid.cell(row=end_row, column=4)
    end_cell.value = (
        f"=IF(MAX({cumulative_range})>={target_ref},_xlfn.MINIFS({dates_range},{cumulative_range},\">=\"&{target_ref}),"
        f"\"Encara falten hores\")"
    )
    end_cell.number_format = "DD/MM/YYYY"
    end_cell.alignment = Alignment(horizontal="left")

    # Print only the calendar: the instructions are for editing the file.
    grid.print_area = f"A1:{get_column_letter(GRID_COLUMNS)}{row}"
    grid.page_setup.orientation = "portrait"
    grid.page_setup.fitToWidth = 1
    grid.page_setup.fitToHeight = 0
    grid.sheet_properties.pageSetUpPr.fitToPage = True

    # Configuration sheet: what produced this calendar.
    config.append(["Horari setmanal"])
    config["A1"].font = Font(bold=True)
    for title, weekday_minutes in _schedule_sections(result, schedule):
        if title:
            config.append([title])
            config.cell(row=config.max_row, column=1).font = Font(italic=True, color="607165")
        for index, weekday in enumerate(("Dilluns", "Dimarts", "Dimecres", "Dijous", "Divendres")):
            minutes = weekday_minutes.get(index, 0)
            config.append([weekday, format_minutes_label(minutes) if minutes else "-"])
    config.append([])
    config.append(["Dies sense pràctiques aplicats"])
    config.cell(row=config.max_row, column=1).font = Font(bold=True)
    config.append(["Inici", "Fi", "Motiu", "Origen"])
    for cell in config[config.max_row]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    if applied_periods:
        for period in applied_periods:
            config.append([period["start"], period["end"], period.get("label") or None, period.get("source")])
    else:
        config.append(["-"])
    config.append([])
    config.append(["Generat amb Tempo el", format_display_date(date.today())])
    for column, width in zip("ABCD", (26, 14, 36, 18)):
        config.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
